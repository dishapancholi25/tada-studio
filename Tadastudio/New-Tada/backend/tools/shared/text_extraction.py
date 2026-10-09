"""Shared text extraction utilities for Microsoft Graph MCP tools.

Extracts readable text from binary file formats (docx, pdf, xlsx, pptx)
downloaded from SharePoint or OneDrive.
"""

import logging
from io import BytesIO
from typing import Any, Dict, Optional, Set

logger = logging.getLogger(__name__)

# Max size for binary text extraction (5MB)
MAX_EXTRACT_SIZE = 5_000_000
# Truncate extracted text to avoid overwhelming LLM context
MAX_TEXT_LENGTH = 100_000

TEXT_EXTENSIONS: Set[str] = {
    ".txt",
    ".csv",
    ".md",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".css",
    ".scss",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".cfg",
    ".conf",
    ".sh",
    ".bash",
    ".ps1",
    ".bat",
    ".cmd",
    ".sql",
    ".r",
    ".rb",
    ".go",
    ".java",
    ".kt",
    ".swift",
    ".log",
    ".env",
    ".gitignore",
    ".dockerfile",
}


def extract_text_from_binary(
    content_bytes: bytes, file_name: str, mime_type: str
) -> Optional[Dict[str, Any]]:
    """Extract readable text from binary file formats.

    Supports docx, pdf, xlsx, and pptx. Returns None for unsupported formats.
    """
    ext = ""
    if "." in file_name:
        ext = f".{file_name.rsplit('.', 1)[-1].lower()}"

    try:
        if (
            ext == ".docx"
            or mime_type
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ):
            return _extract_docx(content_bytes)
        elif ext == ".pdf" or mime_type == "application/pdf":
            return _extract_pdf(content_bytes)
        elif ext in (".xlsx", ".xls") or mime_type in (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel",
        ):
            return _extract_xlsx(content_bytes)
        elif ext in (".pptx", ".ppt") or mime_type in (
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "application/vnd.ms-powerpoint",
        ):
            return _extract_pptx(content_bytes)
    except Exception as e:
        logger.warning("Binary text extraction failed for %s: %s", file_name, e)
        return None

    return None


def _extract_docx(content_bytes: bytes) -> Dict[str, Any]:
    """Extract text from a DOCX file using python-docx."""
    from docx import Document

    doc = Document(BytesIO(content_bytes))
    parts = []

    # Extract paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    # Extract tables
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append(" | ".join(cell.text for cell in row.cells))
        if rows:
            parts.append("\n".join(rows))

    text = "\n\n".join(parts)
    if len(text) > MAX_TEXT_LENGTH:
        text = (
            text[:MAX_TEXT_LENGTH]
            + "\n\n[Content truncated — file too large for full extraction]"
        )

    return {"text": text, "method": "python-docx"}


def _extract_pdf(content_bytes: bytes) -> Dict[str, Any]:
    """Extract text from a PDF file using pypdf."""
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(content_bytes))
    pages = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            pages.append(page_text)

    text = "\n\n".join(pages)
    if len(text) > MAX_TEXT_LENGTH:
        text = (
            text[:MAX_TEXT_LENGTH]
            + "\n\n[Content truncated — file too large for full extraction]"
        )

    return {"text": text, "method": "pypdf", "page_count": len(reader.pages)}


def _extract_xlsx(content_bytes: bytes) -> Dict[str, Any]:
    """Extract text from an Excel file using openpyxl."""
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(content_bytes), read_only=True, data_only=True)
    parts = []

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cell_values = [str(c) if c is not None else "" for c in row]
            if any(v for v in cell_values):
                rows.append("| " + " | ".join(cell_values) + " |")
        if rows:
            # Insert GFM separator after first row (header)
            header = rows[0]
            col_count = header.count("|") - 1
            separator = "|" + " --- |" * col_count
            table_rows = [header, separator] + rows[1:]
            parts.append(f"## Sheet: {sheet_name}\n\n" + "\n".join(table_rows))

    sheet_count = len(wb.sheetnames)
    wb.close()
    text = "\n\n".join(parts)
    if len(text) > MAX_TEXT_LENGTH:
        text = (
            text[:MAX_TEXT_LENGTH]
            + "\n\n[Content truncated — file too large for full extraction]"
        )

    return {"text": text, "method": "openpyxl", "sheet_count": sheet_count}


def _extract_shapes_text(shapes, parts: list) -> None:
    """Recursively extract text from a shapes collection.

    Handles regular shapes with text frames, tables, and group shapes
    containing nested shapes.
    """
    from pptx.shapes.group import GroupShape

    for shape in shapes:
        if isinstance(shape, GroupShape):
            _extract_shapes_text(shape.shapes, parts)
            continue

        if shape.has_table:
            table = shape.table
            rows = []
            for row in table.rows:
                cell_texts = [cell.text for cell in row.cells]
                rows.append(" | ".join(cell_texts))
            if rows:
                parts.append("\n".join(rows))
            continue

        if shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                parts.append(text)


def _extract_pptx(content_bytes: bytes) -> Dict[str, Any]:
    """Extract text from a PowerPoint file using python-pptx."""
    from pptx import Presentation

    prs = Presentation(BytesIO(content_bytes))
    parts = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_parts: list[str] = []
        _extract_shapes_text(slide.shapes, slide_parts)

        if slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            if notes_tf and notes_tf.text.strip():
                slide_parts.append(f"[Notes: {notes_tf.text.strip()}]")

        if slide_parts:
            parts.append(f"--- Slide {slide_num} ---\n" + "\n".join(slide_parts))

    text = "\n\n".join(parts)
    if len(text) > MAX_TEXT_LENGTH:
        text = (
            text[:MAX_TEXT_LENGTH]
            + "\n\n[Content truncated — file too large for full extraction]"
        )

    return {"text": text, "method": "python-pptx", "slide_count": len(prs.slides)}
