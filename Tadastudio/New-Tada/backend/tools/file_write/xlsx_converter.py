"""XLSX generation utilities for the file write tool.

Converts markdown tables or structured JSON to professionally styled
Excel spreadsheets using openpyxl.
"""

import io
import json
import logging
import re
from typing import Any, Dict, List, Optional, Tuple

logger = logging.getLogger(__name__)


# =============================================================================
# TEMPLATE PALETTES
# =============================================================================

TEMPLATE_PALETTES = {
    "modern": {
        "header_fill": "2D2D2D",
        "header_font": "FFFFFF",
        "alt_row": "F6F8FA",
        "border": "E1E4E8",
        "accent": "0066CC",
        "title_fill": "111111",
        "title_font": "FFFFFF",
    },
    "executive": {
        "header_fill": "1A3A5C",
        "header_font": "FFFFFF",
        "alt_row": "F0F4F8",
        "border": "D4DCE6",
        "accent": "1A3A5C",
        "title_fill": "1A3A5C",
        "title_font": "FFFFFF",
    },
    "report": {
        "header_fill": "1E40AF",
        "header_font": "FFFFFF",
        "alt_row": "EFF6FF",
        "border": "BFDBFE",
        "accent": "3B82F6",
        "title_fill": "1E40AF",
        "title_font": "FFFFFF",
    },
    "invoice": {
        "header_fill": "F9FAFB",
        "header_font": "374151",
        "alt_row": "F9FAFB",
        "border": "E5E7EB",
        "accent": "059669",
        "title_fill": "059669",
        "title_font": "FFFFFF",
    },
    "minimal": {
        "header_fill": "FFFFFF",
        "header_font": "333333",
        "alt_row": "FAFAFA",
        "border": "E5E5E5",
        "accent": "333333",
        "title_fill": "333333",
        "title_font": "FFFFFF",
    },
}


def _parse_markdown_tables(
    content: str,
) -> List[Tuple[Optional[str], List[str], List[List[str]]]]:
    """Parse markdown content into tables with optional preceding headings.

    Returns list of (title, headers, rows) tuples.
    """
    lines = content.split("\n")
    tables = []
    current_title = None
    i = 0

    while i < len(lines):
        line = lines[i]

        # Track headings as potential table titles
        heading_match = re.match(r"^#{1,6}\s+(.+)$", line)
        if heading_match:
            current_title = heading_match.group(1).strip()
            i += 1
            continue

        # Detect table start
        if "|" in line and line.strip().startswith("|"):
            table_lines = []
            while (
                i < len(lines) and "|" in lines[i] and lines[i].strip().startswith("|")
            ):
                table_lines.append(lines[i])
                i += 1

            if len(table_lines) >= 2:
                # Parse header
                headers = [
                    c.strip() for c in table_lines[0].strip().strip("|").split("|")
                ]
                # Skip separator line, parse data rows
                rows = []
                for tl in table_lines[2:]:
                    if tl.strip():
                        row = [c.strip() for c in tl.strip().strip("|").split("|")]
                        rows.append(row)
                tables.append((current_title, headers, rows))
                current_title = None
            continue

        i += 1

    return tables


def _parse_structured_json(content: str) -> Optional[Dict[str, Any]]:
    """Try to parse content as structured JSON for spreadsheet generation.

    Expected format:
    {
        "sheets": [
            {
                "name": "Sheet Name",
                "headers": ["Col1", "Col2", ...],
                "rows": [["val1", "val2", ...], ...],
                "column_widths": [15, 20, ...],  // optional
                "number_formats": {"Col2": "#,##0.00"},  // optional
            }
        ],
        "title": "Workbook Title"  // optional
    }
    """
    try:
        data = json.loads(content)
        if isinstance(data, dict) and "sheets" in data:
            return data
        # Support simple format: just headers + rows
        if isinstance(data, dict) and "headers" in data and "rows" in data:
            return {"sheets": [data]}
        # Support list of objects
        if isinstance(data, list) and len(data) > 0 and isinstance(data[0], dict):
            headers = list(data[0].keys())
            rows = [[str(item.get(h, "")) for h in headers] for item in data]
            return {"sheets": [{"name": "Data", "headers": headers, "rows": rows}]}
    except (json.JSONDecodeError, TypeError):
        pass
    return None


def markdown_to_xlsx(
    content: str,
    template: Optional[str] = None,
) -> bytes:
    """Convert markdown tables or structured JSON to a styled XLSX spreadsheet.

    Accepts either:
    - Markdown content containing tables (each table becomes a sheet)
    - JSON with structured data (headers, rows, optional formatting)

    Args:
        content: Markdown with tables or structured JSON.
        template: Template name ('modern', 'executive', 'report', 'invoice', 'minimal').

    Returns:
        XLSX file content as bytes.

    Raises:
        ValueError: If no table data found in content.
        RuntimeError: If XLSX generation fails.
    """
    if not content or not content.strip():
        raise ValueError("Content cannot be empty")

    template_name = template if template in TEMPLATE_PALETTES else "modern"
    palette = TEMPLATE_PALETTES[template_name]

    logger.info(
        "[XLSX-CONVERTER] Converting content to XLSX (template=%s)", template_name
    )

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Try structured JSON first, fall back to markdown tables
        structured = _parse_structured_json(content)

        if structured:
            sheets_data = structured.get("sheets", [])
        else:
            # Parse markdown tables
            tables = _parse_markdown_tables(content)
            if not tables:
                raise ValueError(
                    "No table data found. Provide either markdown with tables "
                    "or structured JSON with 'headers' and 'rows'."
                )
            sheets_data = []
            for title, headers, rows in tables:
                sheets_data.append(
                    {
                        "name": title or f"Sheet {len(sheets_data) + 1}",
                        "headers": headers,
                        "rows": rows,
                    }
                )

        # Style definitions
        header_fill = PatternFill(
            start_color=palette["header_fill"],
            end_color=palette["header_fill"],
            fill_type="solid",
        )
        header_font = Font(
            name="Calibri", bold=True, size=10, color=palette["header_font"]
        )
        alt_row_fill = PatternFill(
            start_color=palette["alt_row"],
            end_color=palette["alt_row"],
            fill_type="solid",
        )
        thin_border = Border(
            left=Side(style="thin", color=palette["border"]),
            right=Side(style="thin", color=palette["border"]),
            top=Side(style="thin", color=palette["border"]),
            bottom=Side(style="thin", color=palette["border"]),
        )
        header_border = Border(
            left=Side(style="thin", color=palette["border"]),
            right=Side(style="thin", color=palette["border"]),
            top=Side(style="medium", color=palette["accent"]),
            bottom=Side(style="medium", color=palette["accent"]),
        )
        data_font = Font(name="Calibri", size=10, color="333333")
        data_alignment = Alignment(vertical="center", wrap_text=True)
        header_alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=True
        )

        for sheet_data in sheets_data:
            sheet_name = str(sheet_data.get("name", "Sheet"))[
                :31
            ]  # Excel 31-char limit
            ws = wb.create_sheet(title=sheet_name)

            headers = sheet_data.get("headers", [])
            rows = sheet_data.get("rows", [])
            column_widths = sheet_data.get("column_widths", None)
            number_formats = sheet_data.get("number_formats", {})

            if not headers:
                continue

            num_cols = len(headers)

            # --- Write headers ---
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.border = header_border
                cell.alignment = header_alignment

            # --- Write data rows ---
            for row_idx, row_data in enumerate(rows, 2):
                for col_idx, value in enumerate(row_data, 1):
                    if col_idx > num_cols:
                        break

                    # Try to convert numeric strings to numbers
                    cell_value = value
                    if isinstance(value, str):
                        # Strip currency symbols and commas for number detection
                        cleaned = re.sub(r"[$,\xA3\u20AC%]", "", value.strip())
                        try:
                            if "." in cleaned:
                                cell_value = float(cleaned)
                            elif cleaned.isdigit() or (
                                cleaned.startswith("-") and cleaned[1:].isdigit()
                            ):
                                cell_value = int(cleaned)
                        except (ValueError, IndexError):
                            pass

                    cell = ws.cell(row=row_idx, column=col_idx, value=cell_value)
                    cell.font = data_font
                    cell.border = thin_border
                    cell.alignment = data_alignment

                    # Apply number format if specified
                    header_name = (
                        headers[col_idx - 1] if col_idx - 1 < len(headers) else ""
                    )
                    if header_name in number_formats:
                        cell.number_format = number_formats[header_name]
                    elif isinstance(cell_value, float):
                        cell.number_format = "#,##0.00"
                    elif "%" in str(value):
                        cell.number_format = "0.0%"

                    # Alternating row colors
                    if row_idx % 2 == 0:
                        cell.fill = alt_row_fill

            # --- Auto-fit column widths ---
            for col_idx in range(1, num_cols + 1):
                if column_widths and col_idx - 1 < len(column_widths):
                    ws.column_dimensions[
                        get_column_letter(col_idx)
                    ].width = column_widths[col_idx - 1]
                else:
                    max_length = (
                        len(str(headers[col_idx - 1]))
                        if col_idx - 1 < len(headers)
                        else 10
                    )
                    for row_data in rows:
                        if col_idx - 1 < len(row_data):
                            max_length = max(
                                max_length, len(str(row_data[col_idx - 1]))
                            )
                    ws.column_dimensions[get_column_letter(col_idx)].width = min(
                        max_length + 4, 50
                    )

            # Freeze header row
            ws.freeze_panes = "A2"

            # Auto-filter
            if rows:
                ws.auto_filter.ref = f"A1:{get_column_letter(num_cols)}{len(rows) + 1}"

        # Save to bytes
        buffer = io.BytesIO()
        wb.save(buffer)
        xlsx_bytes = buffer.getvalue()

        logger.info(
            "[XLSX-CONVERTER] Successfully generated XLSX (%d bytes, %d sheets, template=%s)",
            len(xlsx_bytes),
            len(sheets_data),
            template_name,
        )

        return xlsx_bytes

    except ImportError as e:
        logger.error("[XLSX-CONVERTER] openpyxl not installed: %s", e)
        raise RuntimeError(
            "XLSX generation requires openpyxl. "
            "Please install it with: pip install openpyxl"
        ) from e
    except ValueError:
        raise
    except Exception as e:
        logger.error("[XLSX-CONVERTER] XLSX generation failed: %s", e)
        raise RuntimeError(f"Failed to generate XLSX: {str(e)}") from e
