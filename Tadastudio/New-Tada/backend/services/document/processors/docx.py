"""DOCX file processor using python-docx."""

import logging
from pathlib import Path
from typing import Any, Dict, List

from ..config import DOCX_FILE_TYPES
from ..exceptions import DocumentProcessingError, ProcessorNotAvailableError
from .base import DocumentProcessor


logger = logging.getLogger(__name__)


class DOCXProcessor(DocumentProcessor):
    """Processor for DOCX files using python-docx."""

    def __init__(self):
        """Initialize DOCX processor."""
        self._check_dependencies()

    def _check_dependencies(self):
        """Check if python-docx is available."""
        try:
            import docx  # noqa: F401

            logger.debug("[DOC-DOCX] python-docx available")
        except ImportError:
            raise ProcessorNotAvailableError(
                "docx", "python-docx library not installed"
            )

    def supports(self, file_type: str) -> bool:
        """Check if this processor supports the file type."""
        return file_type in DOCX_FILE_TYPES

    def process(self, file_path: Path, config: Dict[str, Any]) -> Dict[str, Any]:
        """
        Process a DOCX file.

        Args:
            file_path: Path to DOCX file
            config: Processing configuration

        Returns:
            Dictionary with extracted content
        """
        logger.info(f"[DOC-DOCX] Processing DOCX file: {file_path.name}")

        try:
            from docx import Document

            doc = Document(file_path)
            content_parts = []

            # Extract paragraphs
            content_parts.extend(self._extract_paragraphs(doc))

            # Extract tables if configured
            if config.get("extract_tables", True):
                content_parts.extend(self._extract_tables(doc))

            content = "\n\n".join(content_parts)

            logger.debug(f"[DOC-DOCX] Extracted {len(content)} characters")

            return {
                "content": content,
                "extraction_method": "python-docx",
            }

        except Exception as e:
            logger.error(f"[DOC-DOCX] DOCX extraction failed: {e}")
            raise DocumentProcessingError(
                f"Failed to process DOCX: {e}",
                file_path=str(file_path),
                processor="docx",
            )

    def _extract_paragraphs(self, doc) -> List[str]:
        """Extract paragraphs from document."""
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)
        return paragraphs

    def _extract_tables(self, doc) -> List[str]:
        """Extract tables from document."""
        tables = []
        for table in doc.tables:
            table_text = self._format_table(table)
            if table_text:
                tables.append(table_text)
        return tables

    def _format_table(self, table) -> str:
        """Format table as text."""
        table_rows = []
        for row in table.rows:
            row_text = [cell.text for cell in row.cells]
            table_rows.append(" | ".join(row_text))
        return "\n".join(table_rows)
