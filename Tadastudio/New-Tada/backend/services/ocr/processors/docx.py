"""DOCX processor for AI model Vision-based OCR."""

import io
import logging

from docx import Document

from ..base import OCRProcessor
from ..config import SUPPORTED_DOCX_EXTENSIONS, get_default_prompt
from ..exceptions import InvalidInputError
from ..models import OCRRequest, OCRResponse
from ..utils import build_error_response, build_success_response
from .image import ImageProcessor


logger = logging.getLogger(__name__)


class DOCXProcessor(OCRProcessor):
    """Processor for DOCX files with text and image extraction."""

    def __init__(self, client, deployment_name: str, detail: str = "high"):
        """
        Initialize DOCX processor.

        Args:
            client: Azure OpenAI client
            deployment_name: Model deployment name
            detail: Vision API detail level ("low", "auto", or "high")
        """
        super().__init__(client)
        self.deployment_name = deployment_name
        self.image_processor = ImageProcessor(client, deployment_name, detail=detail)

    def supports_file_type(self, file_ext: str) -> bool:
        """Check if this processor supports the given file type."""
        return file_ext in SUPPORTED_DOCX_EXTENSIONS

    def process(self, request: OCRRequest) -> OCRResponse:
        """
        Process DOCX document, extracting text and processing embedded images.

        Args:
            request: OCR request with DOCX and configuration

        Returns:
            OCR response with extracted text and metadata
        """
        prompt = request.prompt or get_default_prompt(request.doc_type)

        try:
            # Open DOCX
            doc = self._open_docx(request)

            # Extract text content
            text_content = []
            text_content.extend(self._extract_paragraphs(doc))
            text_content.extend(self._extract_tables(doc))

            # Process embedded images
            image_results, total_tokens = self._process_embedded_images(doc, prompt)

            # Combine results
            final_text = self._combine_results(text_content, image_results)

            return build_success_response(
                text=final_text,
                extraction_method="docx_mixed",
                tokens_used=total_tokens,
                model=self.deployment_name,
                has_images=len(image_results) > 0,
            )

        except Exception as e:
            logger.error(f"[OCR-DOCX] Error processing DOCX: {e}", exc_info=True)
            return build_error_response(str(e), "docx_mixed")

    def _open_docx(self, request: OCRRequest) -> Document:
        """Open DOCX from path or bytes."""
        if request.file_path:
            return Document(request.file_path)
        elif request.file_bytes:
            return Document(io.BytesIO(request.file_bytes))
        else:
            raise InvalidInputError("Either file_path or file_bytes must be provided")

    def _extract_paragraphs(self, doc: Document) -> list:
        """Extract paragraphs from document."""
        logger.debug("[OCR-DOCX] Extracting paragraphs")
        paragraphs = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                paragraphs.append(paragraph.text)
        return paragraphs

    def _extract_tables(self, doc: Document) -> list:
        """Extract tables from document."""
        logger.debug(f"[OCR-DOCX] Extracting {len(doc.tables)} tables")
        tables = []
        for table in doc.tables:
            table_text = self._format_table(table)
            tables.append(table_text)
        return tables

    def _format_table(self, table) -> str:
        """Format a table as markdown."""
        if not table.rows:
            return ""

        # Build header row
        header_cells = [cell.text for cell in table.rows[0].cells]
        table_text = "\n| " + " | ".join(header_cells) + " |"

        # Build separator row
        table_text += "\n|" + " --- |" * len(header_cells)

        # Build data rows
        for row in table.rows[1:]:
            row_cells = [cell.text for cell in row.cells]
            table_text += "\n| " + " | ".join(row_cells) + " |"

        return table_text

    def _process_embedded_images(self, doc: Document, prompt: str) -> tuple:
        """Process embedded images in document."""
        logger.debug("[OCR-DOCX] Processing embedded images")
        image_results = []
        total_tokens = 0

        rels = doc.part.rels
        for rel in rels.values():
            if "image" in rel.target_ref:
                result = self._process_single_image(rel, prompt)
                if result:
                    image_results.append(result["text"])
                    total_tokens += result["tokens"]

        logger.debug(
            f"[OCR-DOCX] Processed {len(image_results)} embedded images, "
            f"{total_tokens} tokens"
        )
        return image_results, total_tokens

    def _process_single_image(self, rel, prompt: str) -> dict:
        """Process a single embedded image."""
        try:
            image_data = rel.target_part.blob
            image_prompt = f"Embedded image from document. {prompt}"

            image_request = OCRRequest(
                file_bytes=image_data, prompt=image_prompt, doc_type="generic"
            )

            result = self.image_processor.process(image_request)

            if result.success:
                return {"text": result.text, "tokens": result.tokens_used}
            else:
                logger.warning(f"[OCR-DOCX] Failed to process image: {result.error}")
                return None

        except Exception as img_error:
            logger.warning(f"[OCR-DOCX] Failed to process embedded image: {img_error}")
            return None

    def _combine_results(self, text_content: list, image_results: list) -> str:
        """Combine text and image results."""
        final_text = "\n\n".join(text_content)

        if image_results:
            final_text += "\n\n## Embedded Images\n\n" + "\n\n".join(image_results)

        return final_text
